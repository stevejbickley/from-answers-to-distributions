"""Update the supplied Word drafts from audited publication outputs.

Run with the bundled artifact Python, after scripts 04, 09, 05, and 11.
Original native Word equations and author metadata are retained in the main text.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import sys
import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from src.publication_text import MAIN_CAPTIONS, MAIN_FIGURES, SI_FIGURES, revised_results

ROOT=Path(__file__).resolve().parents[1]
R=ROOT/'results'; OUT=ROOT/'manuscript/publication'
MODELS=['gpt4o_anchor','gpt56_sol','jev']
NAMES={'gpt4o_anchor':'GPT-4o','gpt56_sol':'GPT-5.6 Sol','jev':'Jev'}
TITLE='Can AI uncertainty recover human population variation? Token uncertainty, decision uncertainty, and cross-cultural survey responses'


def read(name): return pd.read_csv(R/name)
def fmt(value, digits=3):
    if pd.isna(value):return 'NA'
    v=float(value)
    if abs(v)<.5*10**(-digits):v=0
    return f'{v:.{digits}f}'
def ci(r,stem):return f"{fmt(r[stem])} ({fmt(r[stem+'_ci95_low'])}, {fmt(r[stem+'_ci95_high'])})"

def values():
    d={}
    s=read('tables/table2_primary_summary.csv');s=s[s.representation=='full'].set_index('condition')
    p=read('cultural_map_prompting_summary.csv');p=p[p.representation=='expected'].set_index('condition')
    f=read('tables/table_s5_full_vs_argmax.csv').set_index('condition')
    t=read('tables/table3_population_specificity.csv').set_index('condition')
    for c in MODELS:
        d['js_'+c]=s.loc[c,'js_mean'];d['arg_'+c]=f.loc[c,'mean_js_argmax'];d['diff_'+c]=f.loc[c,'mean_argmax_minus_full'];d['better_'+c]=f.loc[c,'pct_full_better']
        d['map_u_'+c]=p.loc[c,'unconditioned_distance_mean'];d['map_c_'+c]=p.loc[c,'country_conditioned_distance_mean'];d['map_pct_'+c]=p.loc[c,'pct_countries_improved']
        for key,col in [('gain','gain_vs_default_mean'),('gain_lo','gain_vs_default_ci95_low'),('gain_hi','gain_vs_default_ci95_high'),('gain_pct','pct_country_conditioning_improves'),('loco_pct','pct_model_beats_loco_human')]:d[key+'_'+c]=t.loc[c,col]
    contrasts=json.loads((R/'tables/paired_contrasts.json').read_text())
    for k,col in [('mean','mean_error_a_minus_b'),('lo','ci95_low'),('hi','ci95_high')]:d['sj_'+k]=contrasts['gpt56_sol_vs_jev']['js'][col]
    for stem,path,col in [('prompt','prompt_sensitivity.csv','js_to_condition_mean'),('label','openai_label_sensitivity.csv','js_to_label_mean'),('order','option_order_metrics.csv','js_from_first_order')]:
        x=read(path);x=x[x.country!='__DEFAULT__']
        for c,g in x.groupby('condition'):d[stem+'_'+c]=g[col].mean()
    return d


def style_document(doc):
    for name,sz in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Caption',10),('Table Note',9)]:
        if name not in doc.styles:
            doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0)
        st.font.bold=name in ['Title','Heading 1','Heading 2']
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.space_before=Pt(10 if name.startswith('Heading') else 0)
        st.paragraph_format.keep_with_next=name in ['Title','Heading 1','Heading 2']
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    # Remove the bundled Word template's blue title rule and theme font residue.
    for st in doc.styles:
        for tag in ['w:pBdr','w:contextualSpacing']:
            for el in st.element.xpath('.//'+tag):el.getparent().remove(el)
        for fonts in st.element.xpath('.//w:rFonts'):
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+attr),None)
    for sec in doc.sections:
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7)
        sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
        sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
        for part in [sec.header,sec.footer]:
            for p in part.paragraphs:p.clear()
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    for p in doc.paragraphs:
        if p.text.strip():p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for tag in ['w:pBdr','w:contextualSpacing']:
            for el in p._p.xpath('.//'+tag):el.getparent().remove(el)
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.line_spacing=1.12
        # Remove source-template fixed heights/indents, retaining native equation XML.
        p.paragraph_format.left_indent=Pt(0);p.paragraph_format.right_indent=Pt(0)
        if p.style.name not in ['List Paragraph']:
            p.paragraph_format.first_line_indent=Pt(0)
        for run in p.runs:
            run.font.name='Times New Roman';run.font.size=Pt(11)
    return doc


def para(doc,text='',style=None):
    p=doc.add_paragraph(text,style=style);return p

def page(doc):doc.add_page_break()

def table(doc,headers,rows,widths=None,font=9):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [6.9/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for cell,h,w in zip(t.rows[0].cells,headers,widths):cell.text=str(h);cell.width=Inches(w)
    for row in rows:
        cells=t.add_row().cells
        for cell,val,w in zip(cells,row,widths):cell.text=str(val);cell.width=Inches(w)
    for j,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if j==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for cell in row.cells:
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,val in [('top',55),('bottom',55),('left',65),('right',65)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if j==0 or j%2==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF1' if j==0 else 'F5F6F7');tcpr.append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.0
                p.paragraph_format.keep_with_next=False
                for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(font);run.bold=j==0
    return t


def caption(doc,text):return para(doc,text,'Caption')
def note(doc,text):return para(doc,text,'Table Note')

def add_figure(doc,stem,number,legend,maxheight=7.25):
    path=ROOT/'figures'/f'{stem}.png';im=Image.open(path)
    width=min(6.9,maxheight*im.width/im.height)
    p=para(doc);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    # A separate break paragraph can spill onto an otherwise empty page.
    p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',f'Figure {number}. {legend}')
    caption(doc,f'Figure {number}. {legend}')


def remove(p):
    el=p._element;el.getparent().remove(el)


def replace(p,text):
    p.text=text
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11)


def insert_blocks_before(doc,anchor,blocks):
    for kind,text in blocks:
        p=doc.add_paragraph(text,style={'h1':'Heading 1','h2':'Heading 2'}.get(kind,'Normal'))
        anchor.addprevious(p._p)


def repair_source_math(doc):
    """Repair misplaced source delimiter separators without rasterizing equations."""
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    replacements={
        'y∣cq':'y | c, q',
        'y∣θq':'y | θ, q',
        'y∣stateA':'y | state, A',
        'y=k∣y∈Aprompt':'y = k | t ∈ A, prompt',
    }
    repaired=0
    for delimiter in doc.element.xpath('.//m:d'):
        key=''.join(delimiter.xpath('.//m:t/text()',namespaces=ns))
        if key not in replacements:continue
        for child in list(delimiter):delimiter.remove(child)
        props=OxmlElement('m:dPr');delimiter.append(props)
        expression=OxmlElement('m:e');delimiter.append(expression)
        run=OxmlElement('m:r');expression.append(run)
        text=OxmlElement('m:t');text.text=replacements[key];run.append(text)
        repaired+=1
    if repaired!=4:raise AssertionError(f'Unexpected conditional expression count: {repaired}')
    # Literal vertical bars render consistently across Word and LibreOffice.
    for text in doc.element.xpath('.//m:t'):
        if text.text:text.text=text.text.replace('∣','|').replace('∥','‖')
    for delimiter_char in doc.element.xpath('.//m:begChr | .//m:endChr | .//m:sepChr'):
        value=delimiter_char.get(qn('m:val'),'')
        delimiter_char.set(qn('m:val'),value.replace('∣','|').replace('∥','‖'))
    for p in doc.paragraphs:
        if p.text.startswith('probability vectors before aggregation'):
            wrapper=p._p.find(qn('m:oMathPara'));math=wrapper.find(qn('m:oMath'))
            wrapper.remove(math);p._p.replace(wrapper,math)
            for br in math.xpath('.//w:br',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
                br.getparent().remove(br)
    # Source inline equations occasionally abut prose without a space.
    for p in doc.paragraphs:
        for math in p._p.xpath('./m:oMath'):
            for sibling,after in [(math.getprevious(),False),(math.getnext(),True)]:
                if sibling is None or sibling.tag!=qn('w:r'):continue
                texts=sibling.findall(qn('w:t'))
                if not texts:continue
                text=texts[0] if after else texts[-1]
                value=text.text or ''
                if value and (value[0] if after else value[-1]).isalnum():
                    text.text=(' '+value) if after else (value+' ')
                    text.set(qn('xml:space'),'preserve')


def main_manuscript(template,d):
    doc=Document(template);old=list(doc.paragraphs);table1=deepcopy(doc.tables[0]._tbl)
    original_equations=len(doc.element.xpath('.//m:oMath'))
    # Work from stable source paragraphs; remove template content and obsolete figures/tables.
    for p in old[386:]:remove(p)
    body=doc._element.body;start=body.index(old[73]._p);end=body.index(old[149]._p)
    for el in list(body)[start:end]:body.remove(el)
    for p in old[15:26]:
        if p._p.getparent() is not None:remove(p)
    for p in old[:2]:remove(p)
    style_document(doc)
    old[2].style='Title'
    for run in old[2].runs:run.font.size=Pt(17);run.bold=True
    old[26].style='Heading 1';old[26].paragraph_format.page_break_before=True
    old[28].style='Heading 1'
    replace(old[27],f"Human populations contain disagreement, yet a single artificial model is increasingly used to simulate them. Can uncertainty within one model recover variation across many people? Building on Tao et al.’s cross-cultural benchmark, we compare survey-weighted response distributions across 107 countries and territories with GPT-4o and GPT-5.6 Sol next-token probabilities and TypeSafe Jev decision probabilities. Mean Jensen-Shannon divergence is {d['js_gpt4o_anchor']:.3f}, {d['js_gpt56_sol']:.3f}, and {d['js_jev']:.3f}, respectively. Retaining full probabilities substantially improves fidelity relative to the same vectors collapsed to modal answers. Country prompting also improves mean cultural-map location, but improves full response distributions in only {d['gain_pct_gpt4o_anchor']:.1f}%, {d['gain_pct_gpt56_sol']:.1f}%, and {d['gain_pct_jev']:.1f}% of cells. All three models rarely outperform a leave-one-country-out human baseline. Jev most closely matches average human entropy, yet its entropy scarcely covaries with human disagreement. The GPT models have lower mean entropy and positive pooled associations with human entropy, but all country-and-item fixed-effect slope intervals include zero. Model uncertainty therefore contains population-level information without reproducing population heterogeneity. Cultural location, distributional fidelity, and the structure of disagreement are distinct evaluation targets.")
    replace(old[29],'A model probability describes uncertainty within one artificial system; a survey distribution describes differences among people. Across 107 countries and nine cultural-value constructs, retaining full AI probabilities preserves more information than selecting only modal answers. Yet improved cultural-map location does not guarantee better response distributions, and country-conditioned models rarely beat a simple average of other countries’ human responses. Matching the average amount of disagreement also fails to establish that a model identifies where people disagree. Synthetic-population research should assess cultural location, complete response distributions, and the pattern of disagreement separately.')
    for i in [32,33]:remove(old[i])
    old[34].style='Heading 1';old[34].paragraph_format.page_break_before=True
    insert_blocks_before(doc,old[149]._p,revised_results(d))
    old[149].style='Heading 1'
    updates={
    162:'Jev entropy is essentially unrelated to human entropy across pooled country-item cells, and remains weakly associated after item or country-and-item adjustment. The GPT models have lower mean entropy and positive pooled associations, but their two-way fixed-effect intervals include zero. Their response distributions are too concentrated on average; their residual ability to track where disagreement occurs is uncertain.',
    164:'These are different evaluation failures. A model may understate mean disagreement, reproduce its average level while assigning it to the wrong cells, or show an association driven by stable item and country differences. The fixed-effect results limit how strongly the pooled associations can be interpreted as recovery of population-specific heterogeneity.',
    166:'The population-specificity analysis sharpens this interpretation. Country prompting improves mean distributional fidelity relative to each model’s own default, but the improvement is uneven and the GPT-4o interval includes zero. Only 10.9% to 15.2% of cells beat the leave-one-country-out human baseline. Model uncertainty responds to country information without recovering most of the distributional information contained in this simple human comparison. Because the baseline draws on observed human surveys, this is an information benchmark rather than a like-for-like comparison of systems given identical data.',
    168:'The interface analyses show that model uncertainty depends on how it is elicited. Respondent wording, arbitrary labels, and answer ordering alter the recovered distributions. Sensitivity requires repeated usable measurements: a singleton cannot establish stability, and an excluded reference ordering cannot be replaced by a later permutation without changing the estimand. Sol’s lower measured order sensitivity is conditional on selective request retention, including the absence of usable F120 reference comparisons.',
    172:'A stricter residual-mass threshold yields a Sol mean JSD close to the primary estimate, while complete-case restriction leaves only 101 country-item cells and a substantially lower mean JSD. This comparison changes both the probability vectors and the populations and items represented. It therefore does not isolate the effect of the censoring policy or establish that the lower complete-case mean reflects better general fidelity.',
    174:'The cultural-map analysis recovers the broad prompting pattern reported by Tao et al. (1), while changing the human-data release and using probability-derived expected scores. The variant-wise modal reconstruction provides a closer descriptive comparison, but its Y003 marginal treatment and temperature-one probability collection still differ from the original protocol. These results support continuity in the broad pattern rather than exact numerical reproduction.',
    182:'Fourth, finite top-K output bounds rather than reveals the probabilities of omitted labels. Unequal retention affects both aggregation and robustness comparisons. Fifth, model training and post-training histories are incompletely observable, so system differences do not identify an architectural effect. Sixth, Y003 is assessed through marginals rather than its full joint distribution. Seventh, unknown training exposure to survey or cultural data may affect the outputs. Eighth, entropy intervals rely on only nine item clusters, limiting the accuracy of asymptotic cluster-robust inference. Finally, the cross-national human baseline uses observed survey distributions and has a different information source from the model prompts. Frozen identifiers and archived outputs are needed to interpret each result as a model-snapshot comparison.',
    296:'The unconditioned default target supplies the same-model baseline for country-conditioning comparisons and the default cultural-map position. It is excluded from country-level paired inference and the country-conditioned prompt and label sensitivity summaries.',
    318:'Model normalized entropy was regressed on human normalized entropy in pooled, item-fixed-effect, country-fixed-effect, and two-way-fixed-effect specifications. Standard errors use two-way clustering by country and survey item with normal 1.96 critical values. These intervals are approximate because there are only nine item clusters. Undefined covariance estimates are not converted into zero-width intervals. Pearson and Spearman associations, within-item and within-country associations, and model-to-human entropy standard-deviation ratios are reported separately.',
    319:'A unit slope denotes a one-unit fitted change in model entropy per unit of human entropy; a slope near zero indicates little linear association. A slope below one does not imply smaller cross-cell variance, because covariance and variance are distinct quantities. Two-way residual associations remove additive country and item effects by least-squares projection, including in the unbalanced Sol sample.',
    334:'Groups with fewer than two retained descriptors have undefined sensitivity. Summaries exclude the unconditioned target, pool retained descriptor comparisons across country-item cells, and report usable counts. Unequal surviving label counts can give descriptors unequal weight in the aggregated condition mean.',
    338:'Sensitivity is defined only when at least two label assignments survive for the same model, country, item, and descriptor. Singleton groups are retained in diagnostic files with missing sensitivity, rather than being interpreted as zero label dependence. Country-conditioned summaries exclude the unconditioned target and report valid comparison counts.',
    342:'Four answer-order configurations were queried for each model and country-item cell. Configuration 0 is the prespecified reference. Both that request and an alternative request must satisfy the primary censoring policy before their semantic distributions are compared. Cells without a retained reference are excluded from these comparisons; no replacement reference is selected.',
    343:'The design generated 432 requests per model condition. There are 324 valid comparisons for GPT-4o, 157 for Sol, and 324 for Jev. Sol has no usable reference-based F120 comparisons. Mean JSD describes the retained comparisons and does not imply stability of the excluded requests.',
    350:'The projection is repeated with the argmax of each final aggregated vector. A separate variant-wise modal reconstruction first averages retained label assignments within descriptor, selects each item’s modal response, thresholds each Y003 constituent at 0.5, projects complete ten-construct profiles, and averages coordinates over available complete variants. This differs from the source study’s direct point-response protocol. Sol provides 43 country-conditioned profiles under this reconstruction and no complete unconditioned profile; it is excluded from variant-wise modal prompting comparisons.',
    363:'The human benchmark uses the WVS Trend File 1981–2022 version 4.1.0 (11) and EVS Trend File 1981–2017, ZA7503 version 3.0.0 (12). Respondent-level data must be obtained from the respective data providers under their access terms. Analysis code, configuration, aggregate outputs, and reproduction instructions are maintained at https://github.com/stevejbickley/from-answers-to-distributions. The replication package records source hashes, probability-recovery policies, model identifiers, and the dated pricing assumptions used for computational accounting.',
    }
    for i,text in updates.items():replace(old[i],text)
    for p in doc.paragraphs:
        if p.text.startswith('The OpenAI Responses API was queried'):
            replace(p,p.text.replace('for example A and A', 'for example the label A with and without a leading space'))
        if p.text.startswith('Thus, the OpenAI probability object'):
            replace(p,p.text+' In the expression above, t is the returned token and aₖ is the label assigned to response category k.')
    repair_source_math(doc)
    method_blocks=[('h2','Population specificity'),('p','For each observed model-country-item cell, country-conditioning gain equals JSD(human, unconditioned model) minus JSD(human, country-conditioned model). The leave-one-country-out human baseline is the equal-country average of all other countries’ distributions for the same item. Baseline gain is its JSD from the target human distribution minus the model’s JSD. Positive gain always favours the country-conditioned model. Each comparison uses identical cells on both sides. Means and 95% percentile intervals use the crossed country-by-item bootstrap; fractions improved count strictly positive gains and exclude unavailable comparisons.')]
    insert_blocks_before(doc,old[309]._p,method_blocks)
    for i in [197,359,362,364]:old[i].style='Heading 1'
    for p in old[365:384]:
        p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(4)
        for run in p.runs:run.font.size=Pt(10)
    for i in [198,208,217,227,234,241,252,258,265,273,278,293,298,309,321,326,332,336,340,345,352]:
        # Only promote actual short headings; source paragraph indexes may include body text.
        if len(old[i].text)<95 and old[i].text.strip():old[i].style='Heading 2'
    for p in doc.paragraphs:
        if p.text.strip()=='':
            if not p._p.xpath('.//m:oMath | .//w:drawing | .//w:br') and p._p.getparent() is not None: remove(p)
    # Keep display equations with their introducing sentence.
    pars=list(doc.paragraphs)
    for i,p in enumerate(pars):
        if i and not p.text.strip() and p._p.xpath('.//m:oMath'):
            pars[i-1].paragraph_format.keep_with_next=True
    add_main_tables(doc,table1)
    for i,(stem,cap) in enumerate(zip(MAIN_FIGURES,MAIN_CAPTIONS.values()),1):
        add_figure(doc,stem,str(i),cap,maxheight=7.65 if i==1 else 6.6)
    # Confirm preserved native math remains present; removed result captions contained some math.
    remaining=len(doc.element.xpath('.//m:oMath'))
    if remaining<35:raise AssertionError(f'Unexpected equation loss: {original_equations} -> {remaining}')
    path=OUT/'From-answers-to-distributions_Manuscript_20260929_revised.docx';doc.save(path)
    return path,{'source_equations':original_equations,'revised_equations':remaining}


def add_main_tables(doc,table1):
    page(doc);caption(doc,'Table 1. Probability objects compared in the study')
    # Re-create the existing conceptual table in the publication table style.
    from docx.table import Table
    source=Table(table1,doc._body)
    rows=[[c.text for c in row.cells] for row in source.rows]
    table(doc,rows[0],rows[1:],[.72,1.45,1.40,1.60,1.73],font=9)
    note(doc,'Common response support permits numerical comparison without implying identical statistical meaning.')
    page(doc);caption(doc,'Table 2. Primary distributional fidelity')
    s=read('tables/table2_primary_summary.csv');s=s[s.representation=='full'].set_index('condition')
    ct=json.loads((R/'tables/paired_contrasts.json').read_text())['gpt56_sol_vs_jev']
    specs=[('js','js_mean','Jensen-Shannon divergence'),('tv','tv_mean','Total variation'),('wasserstein','wasserstein_mean','Normalized Wasserstein'),('expected_abs_error','expected_abs_error_mean','Absolute expected-score error'),('entropy_abs_error','entropy_abs_error_mean','Absolute normalized-entropy error')]
    rows=[]
    for metric,col,label in specs:
        r=ct[metric];rows.append([label,*[fmt(s.loc[c,col]) for c in MODELS],f"{fmt(r['mean_error_a_minus_b'])} ({fmt(r['ci95_low'],4) if metric=='wasserstein' else fmt(r['ci95_low'])}, {fmt(r['ci95_high'])})",str(r['n_country_item'])])
    table(doc,['Metric',*NAMES.values(),'Sol minus Jev\n(95% CI)','Paired n'],rows,[1.70,.7,.8,.7,1.45,.55])
    note(doc,'Lower errors indicate closer agreement. GPT-4o and Jev means use 963 cells; Sol uses 946. Contrasts are paired on common observed cells, so their means need not equal differences between the displayed marginal means. Wasserstein excludes the unordered trust item. Expected-score error is in original item units and is not normalized across response scales. Intervals use 5,000 crossed bootstrap replicates.')
    page(doc);caption(doc,'Table 3. Cultural prompting and population specificity')
    t=read('tables/table3_country_conditioning.csv').set_index('condition')
    rows=[]
    for c in MODELS:
        r=t.loc[c];rows.append([NAMES[c],f"{fmt(r.map_unconditioned_distance_mean)} → {fmt(r.map_country_conditioned_distance_mean)}",f'{r.map_pct_countries_improved:.1f}%',f'{int(r.map_n_countries)}',f"{fmt(r.distribution_unconditioned_js_mean)} → {fmt(r.distribution_country_conditioned_js_mean)}",f'{r.distribution_pct_cells_improved:.1f}%',str(int(r.n_country_item))])
    table(doc,['Model','Map distance\nwithout → with','Countries\nimproved','Map n','JSD\nwithout → with','Cells\nimproved','Cell n'],rows,[.86,1.30,.85,.5,1.30,.85,.60])
    p=read('tables/table3_population_specificity.csv').set_index('condition')
    rows=[]
    for c in MODELS:
        r=p.loc[c];rows.append([NAMES[c],f'{r.gain_vs_default_mean:.3f} ({r.gain_vs_default_ci95_low:.3f}, {r.gain_vs_default_ci95_high:.3f})',fmt(r.loco_human_js_mean),f'{r.gain_vs_loco_human_mean:.3f} ({r.gain_vs_loco_human_ci95_low:.3f}, {r.gain_vs_loco_human_ci95_high:.3f})',f'{r.pct_model_beats_loco_human:.1f}%'])
    para(doc,'');table(doc,['Model','Gain vs default\n(95% CI)','LOCO\nhuman JSD','Gain vs LOCO\n(95% CI)','Cells beating\nLOCO'],rows,[.9,1.65,1.0,1.75,1.1])
    note(doc,'Map results use expected scores. Distributional gains are baseline JSD minus country-conditioned model JSD; positive values favour country conditioning. Intervals are crossed-bootstrap percentile intervals. Within-model comparisons are paired on identical countries or country-item cells. Model means have different coverage. LOCO denotes the equal-country human distribution for the same item excluding the target country.')


def si_text(doc):
    para(doc,'Supplementary information','Title');para(doc,TITLE,'Heading 1')
    para(doc,'Steven J. Bickley, Ho Fai Chan, Arian Mashhady, Son Tran, and Benno Torgler')
    para(doc,'Corresponding author: Steven J. Bickley\nEmail: s.bickley@qut.edu.au')
    para(doc,'Contents','Heading 1')
    para(doc,'Supplementary methods and results; Tables S1–S16; Figures S1–S11; supplementary references.')
    page(doc)
    blocks=[
    ('Study scope and data provenance','The study compares nine finite-response distributions and four Y003 marginals across 107 countries and territories. Human estimates use WVS Trend v4.1.0 and EVS ZA7503 v3.0.0, common waves 5–7, and years 2005–2022. Within country-years, valid responses are weighted by S017; country-year probability vectors are then averaged equally. The frozen country universe uses S003 identifiers and excludes Egypt, Kuwait, Qatar, Tajikistan, and Uzbekistan. The prepared benchmark contains 392,832 respondent records. This updates the WVS release relative to Tao et al. (1), rather than reproducing the original data release exactly.'),
    ('Probability representations','GPT-4o uses gpt-4o-2024-05-13 and Sol uses gpt-5.6-sol with reasoning effort none. Both expose next-token log probabilities at temperature 1.0 with up to 20 alternatives. Permitted labels are mapped back to substantive response categories before averaging. Nine primary items use three cyclic label assignments per descriptor; each binary Y003 marginal uses two. Jev uses the archived served identifier jev-1.13.0 and returns probabilities over declared Choice alternatives. Its four Y003 decisions are batched within a request. Ten respondent descriptors are evaluated for each country and for an unconditioned target. The primary Jev collection therefore comprises 10,800 requests; each OpenAI collection comprises 37,800. Exact constructs and descriptor wording appear in Tables S1 and S2.'),
    ('Incomplete probability support','An absent OpenAI label is censored, rather than known to have exactly zero probability. The analysis conditions on the observed permitted-label mass and retains incomplete requests only when a conservative upper bound on total omitted permitted-label mass is at most 0.001. Sensitivity policies use 0.0001 or require all permitted labels. Table S3 reports request-level diagnostics, including unconditioned and Y003 requests; Table S4 reports country-item outcomes after each policy. Different policies change both surviving requests and cell coverage, so their marginal means are not paired estimates of a pure policy effect.'),
    ('Distributional comparisons and uncertainty','JSD uses base-2 logarithms and is bounded by zero and one. Total variation is half the sum of absolute category-probability differences. Ordered-item Wasserstein distance is divided by the full response-scale range; it is unavailable for unordered trust (A165). Expected-score error remains in the item’s original units. Normalized entropy is Shannon entropy divided by log2 of category count. The argmax comparator assigns mass one to the first highest-probability substantive category in configured order, using exactly the same aggregated probability vector. Paired model contrasts and mean country-conditioning gains use a crossed bootstrap with 5,000 replicates and seed 20260923, independently resampling countries and items. Wilcoxon tests are secondary cell-level summaries and do not account for the crossed dependence structure.'),
    ('Country conditioning and the human baseline','For target country c and item j, default gain is JSD(Hcj, Mdefault,j) minus JSD(Hcj, Mcj). The LOCO human distribution is the equal-country mean of all other countries for item j; the target country is omitted before averaging. LOCO gain is JSD(Hcj, Hminus-c,j) minus JSD(Hcj, Mcj). Each gain is paired within cell, and percentages improved count strictly positive gains. The LOCO benchmark averages over the model’s available cells when a model-specific mean is reported. It uses survey data unavailable in the model prompt and is an information benchmark, not a controlled comparison of identical input sources.'),
    ('Entropy structure','Pooled, item-fixed-effect, country-fixed-effect, and two-way-fixed-effect regressions relate model entropy to human entropy. Covariance estimates cluster by country and item; reported 95% intervals use normal 1.96 critical values. There are only nine item clusters, limiting asymptotic inference. Negative variance estimates are treated as undefined, not zero. Fixed-effect failures do not fall back to pooled slopes. Two-way residuals are obtained by least-squares projection on country and item indicators, including for unbalanced panels. Within-item and within-country correlations are summarized through equal-group Fisher-z means as well as medians. Table S10 separates regression slopes, correlations, and standard-deviation ratios; these quantify different features of uncertainty.'),
    ('Prompt label and order sensitivity','Country-conditioned prompt sensitivity compares each retained descriptor-specific vector, after averaging retained label assignments, with its model-country-item mean. At least two descriptors are required. Label sensitivity compares each retained semantic label-vector with its mean within model, country, item, and descriptor, requiring at least two assignments. Singleton groups have unavailable sensitivity and are excluded from summary denominators. The unconditioned target is excluded from both summaries. In the 12-country option-order experiment, only comparisons to prespecified ordering 0 are valid; reference and alternative must both survive censoring. An excluded reference is never replaced. Tables S8, S13, and S14 report valid counts and the coverage audit is retained in results/option_order_coverage.csv. Sol has no reference-based F120 comparisons.'),
    ('Y002 and Y003','Y002 enumerates 12 ordered pairs of distinct goals. Pairs containing goals 1 and 3 are materialist, those containing 2 and 4 are post-materialist, and all mixed pairs map to the middle category. Y003 permits up to five of eleven qualities, giving 1,024 possible subsets including the empty set. Four marginal probabilities suffice for the expected autonomy index: P(independence) + P(determination) - P(religious faith) - P(obedience). This follows by linearity of expectation without assuming independence. Y003 is excluded from the nine-item JSD analysis; its four marginal errors appear in Table S9.'),
    ('Cultural-map projection and reconstruction','Ten expected scores are standardized using human-derived parameters, projected through the weighted pairwise-correlation PCA with varimax rotation, and rescaled as PC1′ = 1.81 PC1 + 0.38 and PC2′ = 1.61 PC2 - 0.01. Human and model positions share this transformation. Expected-score and aggregated-argmax profiles cover 107 countries for GPT-4o and Jev and 91 for Sol. A variant-wise modal reconstruction averages label assignments within descriptor, selects modal item responses, thresholds Y003 marginals at 0.5, projects complete profiles, and then averages their coordinates. It covers only 43 Sol countries and lacks a complete Sol default profile. Tables S15 and S16 distinguish these representations and samples. They do not establish exact replication of Tao et al.’s direct temperature-zero responses. Figure 1 uses the source archive’s S003 region assignments and exact palette; classifications only control display colours.'),
    ('Computational accounting and reproduction','The archived computational ledger comprises 95,186 distinct API requests, 20,713,572 input tokens, 1,215,251 output tokens, and 21,928,823 total tokens, including option-order and exploratory native-Score requests. Estimated cost is US$44.87 under the 25 September 2026 list-price snapshot. Provider-reported dollar charges are unavailable; the estimate is not an invoice. Jev requests are deduplicated by request identifier before counting batched decisions. Rebuilding analyses, figures, tables, and this document from archived data requires no new provider calls. The native-Score collection is exploratory and does not enter the primary Choice-probability comparisons.'),
    ]
    for h,p in blocks:para(doc,h,'Heading 2');para(doc,p)
    para(doc,'Interpretation of supplementary tables','Heading 2')
    para(doc,'Counts refer to the unit named in each table. NA means undefined or unavailable, not zero. JSD, probability errors, and normalized entropy are dimensionless; cultural-map distance uses the rescaled coordinate units. Results are descriptive unless an interval is explicitly provided. The complete machine-readable columns, request coverage, country-level entropy diagnostics, and analysis provenance remain in the accompanying CSV and JSON files. The formatted tables below emphasize reported estimands and preserve all model, item, country, and representation groups.')


def si_table(doc,num,title,headers,rows,widths=None,note_text='',font=9):
    page(doc);caption(doc,f'Table S{num}. {title}')
    table(doc,headers,rows,widths,font)
    if note_text:note(doc,note_text)


def supplementary_tables(doc):
    s=read('tables/table_s1_survey_constructs.csv')
    si_table(doc,1,'Survey constructs and response representation',['Item','Construct','Human support','Probability representation'],
      [[r.item,r.construct,r.human_support,r.probability_representation] for r in s.itertuples()],[.55,1.8,1.45,3.1],note_text='The nine finite-response items enter the primary JSD analysis. Y003 is analysed through four marginals.')
    para(doc,'Complete source question wording','Heading 2')
    for r in s.itertuples():para(doc,f'{r.item}  {r.construct}','Heading 2');para(doc,r.source_prompt)
    s=read('tables/table_s2_prompt_variants.csv')
    si_table(doc,2,'Ten respondent-descriptor variants',['ID','Unconditioned wording','Country-conditioned wording'],
      [[r.id,r.descriptor,r.country_descriptor] for r in s.itertuples()],[.4,2.75,3.75],note_text='The literal {country} field is replaced with the frozen country or territory name.')
    s=read('tables/table_s3_openai_diagnostics.csv')
    si_table(doc,3,'OpenAI request-level probability diagnostics',['Model','Calls','Complete','Incomplete\n(%)','Primary\nretained (%)','Strict\nretained (%)','Complete\nonly (%)'],
      [[NAMES[r.condition],r.calls,r.complete_calls,fmt(r.incomplete_percent,1),fmt(r.primary_retained_percent,1),fmt(r.strict_retained_percent,1),fmt(r.complete_only_percent,1)] for r in s.itertuples()], [.9,.8,.8,1.0,1.1,1.1,1.1])
    para(doc,'');table(doc,['Model','Mean allowed mass','Median allowed mass','Maximum omitted-mass bound','Invalid generated label (%)'],
      [[NAMES[r.condition],fmt(r.allowed_mass_mean,6),fmt(r.allowed_mass_median,6),fmt(r.missing_mass_upper_bound_max,6),fmt(r.invalid_generated_percent,2)] for r in s.itertuples()],[1.1,1.3,1.3,1.65,1.55])
    note(doc,'All primary requests, including the default target and Y003 constituents. Allowed mass is measured before conditional normalization. The maximum bound includes excluded requests.')
    s=read('tables/table_s4_openai_censoring_sensitivity.csv')
    si_table(doc,4,'Sensitivity to the OpenAI censoring policy',['Model','Policy','Cells','Mean JSD','Mean TV','Expected-score error','Entropy error'],
      [[NAMES[r.condition],r.censoring_policy.replace('_',' '),r.n_country_item,fmt(r.js_mean),fmt(r.tv_mean),fmt(r.expected_abs_error_mean),fmt(r.entropy_abs_error_mean)] for r in s.itertuples()], [.9,1.25,.55,.8,.8,1.2,1.15],note_text='Means use each policy’s retained cells. Coverage and probability aggregation both change across policies; these are not paired policy effects.')
    s=read('tables/table_s5_full_vs_argmax.csv')
    si_table(doc,5,'Paired full-distribution and argmax fidelity',['Model','Paired cells','Full JSD','Argmax JSD','Mean gain','Median gain','Full better (%)'],
      [[NAMES[r.condition],r.n,fmt(r.mean_js_full),fmt(r.mean_js_argmax),fmt(r.mean_argmax_minus_full),fmt(r.median_argmax_minus_full),fmt(r.pct_full_better,1)] for r in s.itertuples()],[1.05,.8,.85,.9,.85,.9,1.05],note_text='Gain is argmax JSD minus full-distribution JSD. Values are descriptive paired summaries.')
    s=read('tables/table_s6_item_summary.csv')
    si_table(doc,6,'Distributional metrics by item',['Item','Model','n','JSD','TV','Wasserstein','Expected-score error','Entropy error'],
      [[r.item,NAMES[r.condition],r.n,fmt(r.js_mean),fmt(r.tv_mean),fmt(r.wasserstein_mean),fmt(r.expected_error),fmt(r.entropy_error)] for r in s.itertuples()],[.5,1,.4,.7,.7,1.1,1.3,1.2],font=8.5,note_text='n is the number of countries. Wasserstein is unavailable for unordered A165. Expected-score error uses original item units.')
    s=read('tables/table_s7_country_summary.csv')
    si_table(doc,7,'Distributional metrics for every country and model',['Country or territory','Model','Items','Mean JSD','Mean TV','Entropy error'],
      [[r.country,NAMES[r.condition],r.n,fmt(r.js_mean),fmt(r.tv_mean),fmt(r.entropy_error)] for r in s.itertuples()],[2.0,1.15,.5,1.05,1.05,1.15],font=8.5,note_text='All 107 countries and three models are retained. Counts are observed primary items, with a maximum of nine.')
    s=read('tables/table_s8_prompt_sensitivity.csv')
    si_table(doc,8,'Respondent-wording sensitivity by item',['Item','Model','Valid variants','Retained variants','Mean JSD','Median JSD'],
      [[r.item,NAMES[r.condition],r.n,r.n_retained,fmt(r.mean_js),fmt(r.median_js)] for r in s.itertuples()],[.6,1.2,1.3,1.4,1.15,1.25],font=8.5,note_text='Country-conditioned targets only. Singleton descriptor groups contribute to retained counts but not valid sensitivity counts. JSD is measured from the model-country-item mean.')
    s=read('tables/table_s9_y003.csv')
    si_table(doc,9,'Y003 marginal probability errors',['Model','Constituent','Countries','Mean absolute error','Mean squared error'],
      [[NAMES[r.condition],r.quality,r.n,fmt(r.mae),fmt(r.mse)] for r in s.itertuples()],[1.0,2.35,.7,1.4,1.45],note_text='Probabilities refer to selection of the individual quality. No joint-selection distribution is inferred.')
    s=read('tables/table_s10_entropy_structure.csv').set_index('condition')
    labels={'n':'Country-item cells','pearson_overall':'Pooled Pearson r','spearman_overall':'Pooled Spearman rho','human_entropy_sd':'Human entropy SD','model_entropy_sd':'Model entropy SD','sd_ratio_model_to_human':'Model / human SD','two_way_resid_pearson':'Two-way residual Pearson r','two_way_resid_spearman':'Two-way residual Spearman rho','two_way_resid_human_sd':'Human residual SD','two_way_resid_model_sd':'Model residual SD','two_way_resid_sd_ratio':'Model / human residual SD','within_item_pearson_fisher_mean':'Within-item Pearson Fisher mean','within_item_spearman_fisher_mean':'Within-item Spearman Fisher mean','within_item_pearson_median':'Within-item Pearson median','within_item_spearman_median':'Within-item Spearman median','within_country_pearson_fisher_mean':'Within-country Pearson Fisher mean','within_country_spearman_fisher_mean':'Within-country Spearman Fisher mean','within_country_pearson_median':'Within-country Pearson median','within_country_spearman_median':'Within-country Spearman median'}
    rows=[[lab,*[str(int(s.loc[c,k])) if k=='n' else fmt(s.loc[c,k]) for c in MODELS]] for k,lab in labels.items()]
    si_table(doc,10,'Entropy level variation and association',['Diagnostic',*NAMES.values()],rows,[3.3,1.2,1.2,1.2],note_text='Fisher means give equal weight to available group correlations after Fisher-z transformation. SD is the standard deviation across observed cells, not the mean entropy within distributions.')
    para(doc,'Regression specifications','Heading 2')
    rows=[]
    for prefix,label in [('ols','Pooled'),('item_fe','Item FE'),('country_fe','Country FE'),('two_way_fe','Country + item FE')]:
        for c in MODELS:
            r=s.loc[c];rows.append([label,NAMES[c],fmt(r[prefix+'_slope']),f"{fmt(r[prefix+'_ci_low'])}, {fmt(r[prefix+'_ci_high'])}",fmt(r[prefix+'_se']),f"{r[prefix+'_p']:.4g}",fmt(r[prefix+'_r2'])])
    table(doc,['Specification','Model','Slope','95% CI','SE','p','R²'],rows,[1.5,1.15,.65,1.4,.65,.65,.6],font=8.5)
    note(doc,'FE denotes fixed effects. Covariance clusters by country and item; normal-approximation intervals use only nine item clusters. Slopes and p values are descriptive evidence subject to that limitation.')
    s=read('tables/table_s11_entropy_by_item.csv')
    si_table(doc,11,'Within-item entropy associations',['Item','Model','n','Pearson r','Spearman rho','Slope','Human mean','Model mean'],
      [[r.item,NAMES[r.condition],r.n,fmt(r.pearson),fmt(r.spearman),fmt(r.slope),fmt(r.human_entropy_mean),fmt(r.model_entropy_mean)] for r in s.itertuples()],[.5,1.1,.4,.9,.95,.75,1.15,1.15],font=8.5,note_text='Associations vary across countries within each item; slopes are descriptive and have no interval here.')
    s=read('tables/table_s12_population_specificity_by_item.csv')
    si_table(doc,12,'Country-conditioning gains by item',['Item','Model','n','Country JSD','Default JSD','Default gain','LOCO gain','Model shift'],
      [[r.item,NAMES[r.condition],r.n,fmt(r.country_model_js_mean),fmt(r.default_model_js_mean),fmt(r.gain_vs_default_mean),fmt(r.gain_vs_loco_human_mean),fmt(r.model_shift_from_default_js_mean)] for r in s.itertuples()],[.5,1.1,.4,.95,.95,.95,.95,1.0],font=8.5,note_text='Gains are baseline minus model JSD; positive favours country conditioning. Model shift is JSD between conditioned and default model distributions. The corresponding LOCO human mean equals country JSD plus LOCO gain; complete values are retained in the CSV.')
    s=read('tables/table_s13_label_sensitivity.csv')
    si_table(doc,13,'OpenAI semantic label-assignment sensitivity',['Model','Item','Valid comparisons','Retained vectors','Mean JSD','Median JSD'],
      [[NAMES[r.condition],r.item,r.n,r.n_retained,fmt(r.mean_js),fmt(r.median_js)] for r in s.itertuples()],[1.2,.6,1.4,1.4,1.1,1.2],note_text='At least two assignments must survive within a model-country-item-descriptor group. Singleton vectors are counted as retained but have undefined sensitivity. Default targets are excluded.')
    s=read('tables/table_s14_option_order_sensitivity.csv')
    rows=[]
    for c in MODELS:
        for item in read('tables/table_s1_survey_constructs.csv').item:
            if item=='Y003':continue
            g=s[(s.condition==c)&(s.item==item)]
            rows.append([NAMES[c],item,int(g.n.iloc[0]) if len(g) else 0,fmt(g.mean_js.iloc[0]) if len(g) else 'NA',fmt(g.median_js.iloc[0]) if len(g) else 'NA'])
    si_table(doc,14,'Sensitivity to answer-option order',['Model','Item','Valid comparisons','Mean JSD','Median JSD'],rows,[1.3,.7,1.8,1.5,1.6],note_text='Reference is always order 0. Both requests must be retained. Up to 36 comparisons per item are available across 12 countries. No Sol F120 comparison has a valid reference; its sensitivity is unavailable.')
    s=read('tables/table_s15_cultural_map.csv')
    reps={'expected':'Expected score','argmax':'Aggregated argmax','tao_modal':'Variant-wise modal'}
    si_table(doc,15,'Cultural-map distance by representation',['Model','Representation','Countries','Mean distance','Median distance'],
      [[NAMES[r.condition],reps[r.representation],r.n,fmt(r.mean_distance),fmt(r.median_distance)] for r in s.itertuples()],[1.2,2.0,.7,1.5,1.5],note_text='Distances compare country-conditioned model and human coordinates. Samples differ by representation, especially for Sol’s variant-wise modal reconstruction.')
    para(doc,'Expected-score and argmax results on the common 91-country subset','Heading 2')
    s=read('cultural_map_common_country_summary.csv');s=s[s.representation.isin(['expected','argmax'])]
    table(doc,['Model','Representation','Countries','Mean distance','Median distance'],[[NAMES[r.condition],reps[r.representation],r.n,fmt(r.mean_distance),fmt(r.median_distance)] for r in s.itertuples()],[1.2,2.0,.7,1.5,1.5])
    s=read('tables/table_s16_tao_replication.csv')
    rows=[]
    for r in s.itertuples():
        lab='Tao GPT-4o (1)' if r.condition=='gpt4o_published' else NAMES[r.condition]
        rows.append([lab,reps.get(r.representation,'Published point'),r.n_countries,fmt(r.unconditioned_distance_mean),fmt(r.country_conditioned_distance_mean),fmt(r.mean_distance_improvement),fmt(r.pct_countries_improved,1)])
    si_table(doc,16,'Cultural prompting and the Tao benchmark',['Source / model','Representation','Countries','Without','With','Mean gain','Improved (%)'],rows,[1.2,1.55,.6,.9,.9,.9,.85],font=8.5,note_text='The external source row reports published point-response results from Tao et al. (1); remaining rows are current-study estimates. Within each row the country set is paired. Variant-wise modal uses marginal Y003 reconstruction and temperature-one probabilities, so it is not an exact source replication. Sol has no complete default variant-wise modal profile and therefore no such prompting row.')


def supplementary(template):
    # Retain source authorship/title context but replace the unfilled journal template.
    source=Document(template)
    if not any('Steven J. Bickley' in p.text for p in source.paragraphs):raise ValueError('Unexpected SI source')
    doc=style_document(Document())
    si_text(doc);supplementary_tables(doc)
    for i,(stem,legend) in enumerate(SI_FIGURES,1):add_figure(doc,stem,f'S{i}',legend,maxheight=6.8)
    page(doc);para(doc,'Supplementary references','Heading 1')
    refs=[
      '1. Tao Y, Viberg O, Baker RS, Kizilcec RF. Cultural bias and cultural alignment of large language models. PNAS Nexus. 2024;3:pgae346. doi:10.1093/pnasnexus/pgae346. Source analysis and cultural-region lookup: https://osf.io/7sj3w/.',
      '2. Haerpfer C, Inglehart R, Moreno A, et al., editors. World Values Survey Trend File 1981–2022. Version 4.1.0. JD Systems Institute and WVSA Secretariat. doi:10.14281/18241.27.',
      '3. European Values Study. EVS Trend File 1981–2017, ZA7503. Version 3.0.0. GESIS Data Archive. doi:10.4232/1.14021.',
    ]
    for text in refs:para(doc,text)
    path=OUT/'From-answers-to-distributions_SI_20260929_revised.docx';doc.save(path);return path


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--main-template',type=Path,default=ROOT/'manuscript/templates/Manuscript_20260929_source.docx')
    parser.add_argument('--si-template',type=Path,default=ROOT/'manuscript/templates/SI_20260929_source.docx')
    args=parser.parse_args();OUT.mkdir(parents=True,exist_ok=True)
    # Some interpretive sentences contain frozen-run estimates. Fail explicitly
    # if their sources change, rather than emitting a partly stale manuscript.
    frozen=json.loads((ROOT/'manuscript/publication_text_sources.json').read_text())
    changed=[name for name,digest in frozen.items() if hashlib.sha256((R/name).read_bytes()).hexdigest()!=digest]
    if changed:
        raise ValueError('Review publication_text.py for changed empirical sources before rebuilding: '+', '.join(changed))
    audit=json.loads((R/'publication_audit.json').read_text())
    if not audit['passed']:raise ValueError('Publication audit has not passed')
    d=values();mainpath,math=main_manuscript(args.main_template,d);sipath=supplementary(args.si_template)
    manifest={'inputs':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [args.main_template,args.si_template]},'results_audit_sha256':hashlib.sha256((R/'publication_audit.json').read_bytes()).hexdigest(),
      'outputs':[str(mainpath.relative_to(ROOT)),str(sipath.relative_to(ROOT))],**math,'values':d}
    (OUT/'document_build_manifest.json').write_text(json.dumps(manifest,indent=2))
    # Editable text companions support review without Word; equations remain native in DOCX.
    for path in [mainpath,sipath]:
        doc=Document(path);lines=['Text-only review companion. Equations, tables and embedded figures are omitted here; the DOCX is the complete manuscript.']
        for p in doc.paragraphs:
            if p.text:lines.append(('# ' if p.style.name=='Title' else '## ' if p.style.name=='Heading 1' else '### ' if p.style.name=='Heading 2' else '')+p.text)
        (OUT/(path.stem+'.md')).write_text('\n\n'.join(lines)+'\n')
    print(mainpath);print(sipath);print(math)

if __name__=='__main__':main()
